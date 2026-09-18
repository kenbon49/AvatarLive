"""Budget three current cloud APIs and a minimal video-playback streaming PC."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT = Path(__file__).resolve().parents[1] / "artifacts" / "AvatarLive-cloud-API-streaming-budget-i7-64GB-2026-09-16.xlsx"
VIDEO_PRICING = "https://help.aliyun.com/zh/avatar/avatar-application/product-overview/avatar-video-pricing"
TTS_PRICING = "https://help.aliyun.com/zh/model-studio/model-pricing"
DEEPSEEK_PRICING = "https://api-docs.deepseek.com/zh-cn/quick_start/pricing/"
COLOR = "193F42"
SHADE = "EAF3EF"
WARNING = "FFF1D8"
WHITE = "FFFFFF"
CNY = '"¥"#,##0.00;[Red]("¥"#,##0.00)'

# Four core components, assuming a compatible chassis, PSU, cooler and peripherals are reused.
HARDWARE = [
    ("Core i7-14700（带核显，非 F）", "20 核 28 线程、UHD 770/Quick Sync；需兼容 14 代 BIOS 和足够散热", 1, 2600, 3500, "浏览器、业务服务、多窗口、OBS/FFmpeg、SRS；保留核显硬编能力"),
    ("B760 DDR4 主板（千兆网口）", "LGA1700；DDR4、双内存槽以上；足够的 CPU 供电/VRM 散热；视频输出和 14 代 BIOS", 1, 900, 1500, "连接 i7 核显、64GB 内存、有线直播网络及显示设备"),
    ("DDR4 64GB 内存", "2 x 32GB 双通道；核实主板 QVL、频率与稳定性", 1, 900, 1600, "前端、多窗口、SRS/FFmpeg、缓存与本机业务服务"),
    ("NVMe SSD 4TB（唯一硬盘）", "系统、视频素材、缓存和录播文件共用", 1, 2000, 3500, "播放/录制素材；无独立备份盘"),
]


def title(ws, headline, subtitle, cols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    cell = ws.cell(1, 1, headline)
    cell.fill = PatternFill("solid", fgColor=COLOR)
    cell.font = Font(name="Microsoft YaHei", size=16, bold=True, color=WHITE)
    ws.row_dimensions[1].height = 39
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    ws.cell(2, 1, subtitle).alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 45


def header(ws, labels):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(4, col, label)
        cell.fill = PatternFill("solid", fgColor=COLOR)
        cell.font = Font(name="Microsoft YaHei", color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 36


def row(ws, index, values, caution=False):
    for col, value in enumerate(values, 1):
        cell = ws.cell(index, col, value)
        cell.font = Font(name="Microsoft YaHei", size=10, color=COLOR)
        cell.fill = PatternFill("solid", fgColor=WARNING if caution else SHADE if index % 2 == 0 else WHITE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[index].height = 50


def widths(ws, sizes):
    for col, size in enumerate(sizes, 1):
        ws.column_dimensions[get_column_letter(col)].width = size


def build():
    wb = Workbook()
    overview = wb.active
    overview.title = "报价总览"
    prices = wb.create_sheet("三项API价格")
    examples = wb.create_sheet("视频时长测算")
    machine = wb.create_sheet("本机直播硬件")
    inputs = wb.create_sheet("计费参数与依据")

    title(inputs, "可修改的计费参数", "B 列是预算输入；高峰按缓存未命中计算。DeepSeek 经项目 LiteLLM 代理时，以代理账单为最终价格。", 4)
    header(inputs, ["参数", "数值", "口径", "来源"])
    settings = [
        ("灵眸数字人视频 CNY/秒", 0.10, "官网 API 后付费 6 元/分钟、精确到秒；非控制台积分套餐", VIDEO_PRICING),
        ("Qwen3-TTS-VC CNY/万字符", 0.80, "单独发起语音合成才计入；不重复计算灵眸成片所用声音", TTS_PRICING),
        ("DeepSeek Pro 高峰输入 CNY/百万 token", 9.0, "缓存未命中，高峰北京工作日 9-12、14-18", DEEPSEEK_PRICING),
        ("DeepSeek Pro 高峰输出 CNY/百万 token", 27.0, "缓存未命中，输出含实际计费 token", DEEPSEEK_PRICING),
        ("DeepSeek Pro 空闲输入 CNY/百万 token", 4.5, "缓存未命中；缓存命中另有更低价", DEEPSEEK_PRICING),
        ("DeepSeek Pro 空闲输出 CNY/百万 token", 13.5, "工作日高峰以外均为空闲时段", DEEPSEEK_PRICING),
        ("每分钟口播字符", 240, "示例字数，可依真实脚本修改", "业务示例假设"),
        ("每分钟 LLM 输入 token", 1000, "含系统提示及材料，视频分钟与 token 无必然对应关系", "业务示例假设"),
        ("每分钟 LLM 输出 token", 500, "含实际计费输出；请按接口 usage 校准", "业务示例假设"),
        ("每分钟 LLM 调用次数", 1, "示例为每分钟一次文本生成，可改成真实调用量", "业务示例假设"),
        ("每月成片分钟", 100, "100 条各 1 分钟或等量时长；月费示例", "业务示例假设"),
        ("额外 TTS 的比例", 0, "0 表示灵眸成片不另调 TTS；1 表示全部另行合成声音", "业务示例假设"),
    ]
    for idx, values in enumerate(settings, 5):
        row(inputs, idx, values)
        if str(values[3]).startswith("https://"):
            inputs.cell(idx, 4).hyperlink = values[3]
            inputs.cell(idx, 4).style = "Hyperlink"
    widths(inputs, [48, 18, 90, 87])
    inputs.freeze_panes = "B5"
    inputs.cell(16, 2).number_format = "0%"

    title(prices, "只计三类云 API", "视频合成、单独的语音合成及 DeepSeek 文本接口。下方价格为 2026-09-16 查阅的公开价，代理/优惠以账单为准。", 7)
    header(prices, ["服务", "项目调用型号", "计费单位", "高峰/普通价", "空闲价", "计费说明", "官方来源"])
    api_rows = [
        ("数字人视频合成 API", "阿里云灵眸 / CreateBroadcastVideoFromTemplate", "视频分钟", "='计费参数与依据'!B5*60", "='计费参数与依据'!B5*60", "¥6/分钟；按实际生成秒数结算。使用公共数字人+音色的模板口播。", VIDEO_PRICING),
        ("语音合成 API", "阿里云 qwen3-tts-vc-2026-01-22", "万字符", "='计费参数与依据'!B6", "='计费参数与依据'!B6", "仅对单独调用接口的语音收费；已用灵眸自带音色成片时无需重复计入。", TTS_PRICING),
        ("大模型 DeepSeek API：输入", "LiteLLM / deepseek-v4-pro", "百万输入 token", "='计费参数与依据'!B7", "='计费参数与依据'!B9", "官网缓存未命中价；项目需选择 llm-deepseek，当前默认仍为 llm-gpt。", DEEPSEEK_PRICING),
        ("大模型 DeepSeek API：输出", "LiteLLM / deepseek-v4-pro", "百万输出 token", "='计费参数与依据'!B8", "='计费参数与依据'!B10", "公开直连单价；若经内部代理结算，需核对代理价。", DEEPSEEK_PRICING),
    ]
    for idx, values in enumerate(api_rows, 5):
        row(prices, idx, values)
        for col in (4, 5):
            prices.cell(idx, col).number_format = CNY
        prices.cell(idx, 7).hyperlink = values[6]
        prices.cell(idx, 7).style = "Hyperlink"
    widths(prices, [29, 53, 24, 21, 19, 85, 84])
    prices.freeze_panes = "D5"

    title(examples, "视频时长费用示例", "示例每分钟 240 字、LLM 输入 1000/output 500 token，缓存未命中；额外 TTS 列仅在另行调用时叠加。", 9)
    header(examples, ["视频分钟", "数字人成片", "单独 TTS（可选）", "DeepSeek 高峰", "常规合计/高峰", "另加 TTS/高峰", "DeepSeek 空闲", "常规合计/空闲", "另加 TTS/空闲"])
    for idx, duration in enumerate((1, 5, 10, "='计费参数与依据'!B15"), 5):
        row(examples, idx, (duration,
            f"=A{idx}*60*'计费参数与依据'!B5",
            f"=A{idx}*'计费参数与依据'!B11/10000*'计费参数与依据'!B6",
            f"=A{idx}*'计费参数与依据'!B14*('计费参数与依据'!B12*'计费参数与依据'!B7+'计费参数与依据'!B13*'计费参数与依据'!B8)/1000000",
            f"=B{idx}+D{idx}", f"=E{idx}+C{idx}",
            f"=A{idx}*'计费参数与依据'!B14*('计费参数与依据'!B12*'计费参数与依据'!B9+'计费参数与依据'!B13*'计费参数与依据'!B10)/1000000",
            f"=B{idx}+G{idx}", f"=H{idx}+C{idx}"), caution=idx == 8)
        for col in range(2, 10):
            examples.cell(idx, col).number_format = CNY
    row(examples, 10, ("月度实际预估", "='计费参数与依据'!B15*60*'计费参数与依据'!B5",
        "='计费参数与依据'!B15*'计费参数与依据'!B11/10000*'计费参数与依据'!B6*'计费参数与依据'!B16",
        "=D8", "=B10+D10", "=E10+C10", "=G8", "=B10+G10", "=H10+C10"))
    for col in range(2, 10):
        examples.cell(10, col).number_format = CNY
    widths(examples, [24, 21, 25, 22, 24, 25, 22, 24, 26])
    examples.freeze_panes = "B5"

    title(machine, "本机直播 | 四项核心硬件预算", "仅包含指定的 CPU、主板、64GB 内存及单块 4TB 盘；依赖复用现有机箱、合适散热器、电源、显示器和有线网络，不是可直接开机的整机价。", 8)
    header(machine, ["硬件", "规格/验收", "数量", "单价低", "单价高", "小计低", "小计高", "用途"])
    for idx, (name, spec, qty, low, high, purpose) in enumerate(HARDWARE, 5):
        row(machine, idx, (name, spec, qty, low, high, f"=C{idx}*D{idx}", f"=C{idx}*E{idx}", purpose), caution=qty == 0)
        for col in (4, 5, 6, 7):
            machine.cell(idx, col).number_format = CNY
    last = 4 + len(HARDWARE)
    row(machine, last + 1, ("四项核心硬件合计", "非整机报价；不含实时 MuseTalk GPU", "", "", "", f"=SUM(F5:F{last})", f"=SUM(G5:G{last})", "采购前核查 BIOS、散热器、电源和现有机箱"))
    machine.cell(last + 1, 6).number_format = CNY
    machine.cell(last + 1, 7).number_format = CNY
    machine.auto_filter.ref = f"A4:H{last}"
    machine.freeze_panes = "C5"
    widths(machine, [38, 66, 11, 18, 18, 18, 18, 67])

    title(overview, "AvatarLive | 三项云 API + 本机直播设备", f"{date(2026, 9, 16).isoformat()}  |  云端生成成片、在本机播放/转推；不采购 H100，也不运行本地实时数字人生成。", 5)
    header(overview, ["方案", "预算低/空闲", "预算高/高峰", "类型", "说明"])
    row(overview, 5, ("四项核心硬件一次性", f"='本机直播硬件'!F{last+1}", f"='本机直播硬件'!G{last+1}", "非整机价", "i7-14700 非 F、B760 DDR4、64GB、单块 4TB；机箱、电源、散热、外设需现有设备复用。"))
    row(overview, 6, ("云 API 月费（示例）", "='视频时长测算'!I10", "='视频时长测算'!F10", "月度：空闲 / 高峰", "默认 100 分钟视频、DeepSeek 每分钟一调用；语音已在灵眸成片时不重复计算。"))
    row(overview, 7, ("单独语音合成（可选）", "='视频时长测算'!C8", None, "月度增量上限", "仅当 100% 成片语音另行单独调用 Qwen3-TTS-VC；实际按参数页比例计算。"))
    for r in (5, 6, 7):
        for col in (2, 3):
            overview.cell(r, col).number_format = CNY
    notes = [
        (9, "1 分钟示例", "灵眸 1 分钟 ¥6；按每分钟输入 1000/输出 500 token，DeepSeek 高峰约 ¥0.0225，合计约 ¥6.02。额外单独合成 240 字 Qwen 音频，再加约 ¥0.0192。"),
        (10, "DeepSeek 切换", "项目目前默认 llm-gpt，预算假设切换到 llm-deepseek/deepseek-v4-pro；不改代码/配置时不会产生这里估算的 DeepSeek 费用。"),
        (11, "计费边界", "灵眸成片按实际生成秒数；Qwen 声音只在额外请求时收费；DeepSeek 受输入/输出/缓存/时段影响。公开价不保证内部 LiteLLM 代理的最终结算。"),
        (12, "直播能力边界", "此方案仅播放云端预生成视频并推流。现有 MuseTalk 实时换嘴、实时互动数字人若保留，需要另配 GPU 和运行服务，不属于本机硬件合计。"),
        (13, "复用设备前提", "四项硬件不能单独开机：须有兼容机箱、可支撑 i7 满载的散热器、品质合格的 650W 左右电源、显示器和有线网络；若没有，须另行采购并加入总价。"),
        (14, "采购与价格", "B760 必须选 DDR4 型号、有足够 VRM 散热、14 代 BIOS 与核显视频输出；更新 BIOS/CPU 微码并按 Intel 默认功率设置测试。价格是立项估算，未询实时卖家，也未含公网/CDN、授权、税、人工和单盘备份。"),
    ]
    for idx, label, body in notes:
        overview.cell(idx, 1, label).fill = PatternFill("solid", fgColor=SHADE)
        overview.cell(idx, 1).value = label
        overview.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=5)
        overview.cell(idx, 2, body).alignment = Alignment(wrap_text=True, vertical="center")
        overview.row_dimensions[idx].height = 46
    widths(overview, [29, 23, 23, 26, 110])

    for ws in wb:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"{OUTPUT}\nhardware_min={sum(q*l for _, _, q, l, h, _ in HARDWARE)} hardware_max={sum(q*h for _, _, q, l, h, _ in HARDWARE)}")


if __name__ == "__main__":
    build()
