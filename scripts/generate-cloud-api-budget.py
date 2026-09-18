"""Create an indicative, editable monthly budget for AvatarLive's cloud APIs."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT = Path(__file__).resolve().parents[1] / "artifacts" / "AvatarLive-cloud-API-monthly-budget-2026-09-16.xlsx"
DAY = date(2026, 9, 16)
DARK = "193F42"
LIGHT = "EAF3EF"
WHITE = "FFFFFF"
AMBER = "FFF1D8"
MONEY = '"¥"#,##0.00;[Red]("¥"#,##0.00)'
PRICE = '"¥"#,##0.000000;[Red]("¥"#,##0.000000)'
ALIYUN_PRICE = "https://help.aliyun.com/zh/model-studio/model-pricing"
AZURE_PRICE = "https://azure.microsoft.com/en-us/pricing/details/speech/"
GOOGLE_PRICE = "https://ai.google.dev/gemini-api/docs/pricing#gemini-3.1-flash-image"
LINGMOU_GUIDE = "https://help.aliyun.com/zh/avatar/avatar-application/user-guide/avatar-video-operation-guide"

# Key, category, model, unit, quantity formula, unit-price formula/value, status, scope, code basis, pricing reference.
LINES = [
    ("lingmou", "视频", "阿里云灵眸透明口播 / CreateBroadcastVideoFromTemplate", "成片分钟（待确认）", "=B17*B18/60", None, "待合同价", "1080p/30fps；计费单位和模板/形象授权以账户合同为准", "src/lib/server/aliyun-avatar-video.ts", LINGMOU_GUIDE),
    ("happyhorse", "视频", "happyhorse-1.0-video-edit / 720P", "输入+输出秒", "=B14*(B12+B13)", 0.90, "公开原价", "同一条任务的输入秒+输出秒；50 条 x (8+8) 秒仅为示例", "src/app/avatar-video-api/jobs/route.ts; services/cost/happyhorse_cost.py", ALIYUN_PRICE),
    ("seedance", "视频", "doubao-seedance-2-5-260628 / 720p", "输出秒（估算）", "=B15*B16", "=B35*B36*B37*B38/1024/1000000", "项目折算价", "项目按 69.91 元/百万 token、1280x720/24fps 折算；实际 -1 自动时长依返回结果", "src/app/avatar-video-api/jobs/route.ts; services/cost/seedance_cost.py", ""),
    ("gemini_image", "图片", "gemini-3.1-flash-image / 1K", "张", "=B19", "=(B6*B9+B7*B10+B8*B11)/1000000*B5", "项目费率估算", "每张假定 2000 输入/1000 文本思考/1120 图片输出 token；以实际 usage 为准", "src/app/avatar-image-api/edit/route.ts; services/cost/gemini_image_cost.py", GOOGLE_PRICE),
    ("qwen_vc", "配音", "qwen3-tts-vc-2026-01-22", "万字符", "=B20/10000", 0.80, "公开价待核区", "合成文本；首次音色注册不重复计费；项目内部成本表仍记 1.0 元/万字", "src/app/live-voice-api/synthesize/route.ts; services/cost/bailian_tts_cost.py", ALIYUN_PRICE),
    ("qwen_vd", "声音设计", "qwen3-tts-vd-2026-01-26", "万字符", "=B21*B22/10000", 0.80, "公开价待核区", "声音设计预览文本；录入 20 次 x 200 字作示例", "src/app/avatar-voice-api/design/route.ts", ALIYUN_PRICE),
    ("azure_tts", "配音可选", "Azure Speech Neural TTS", "百万字符", "=B26/1000000", "=B33*B5", "官网美元换算", "默认用量 0；改走 Azure/MuseTalk TTS 时填字符量；免费额度未抵扣", "apps/api/app/services/tts/azure.py; apps/musetalk/server.py", AZURE_PRICE),
    ("gpt_input", "文本", "LiteLLM -> azure-gpt-5.4 输入", "百万 token", "=B23*B24/1000000", None, "代理价待询", "项目默认 llm-gpt；按代理账户实际输入价/账单填写，不等同公开 OpenAI 价", "apps/api/app/services/llm/config.py; src/app/live-ai-api/expand/route.ts", ""),
    ("gpt_output", "文本", "LiteLLM -> azure-gpt-5.4 输出", "百万 token", "=B23*B25/1000000", None, "代理价待询", "按代理账户实际输出价填写；若有缓存、图像 token、思考 token 另核对", "apps/api/app/services/llm/chat.py", ""),
    ("fish", "声音参考", "Fish Audio 公开预设查询/试听", "次", "=B32", None, "待核算", "仅查询公开样例；克隆走本地 OpenVoice，不把试听当付费合成", "src/lib/server/avatar-voice-service.ts", ""),
    ("gpu", "运行配套", "MuseTalk 实时渲染 GPU 节点", "GPU 台月", "=B27", None, "待询价", "云 API 并未替换本地 MuseTalk；若改租 GPU 需核实规格/并发与云租价", "apps/musetalk/server.py; src/lib/musetalk-total-stream.ts", ""),
    ("app", "运行配套", "Next.js / FastAPI / PostgreSQL / SRS 主机", "台月", "=B28", None, "待询价", "可沿用现有主机；若上云需另计 CPU、内存、公网及数据库", "apps/api/app/main.py; infra/", ""),
    ("oss", "存储", "视频素材/成片对象存储", "GB 月", "=B29", None, "待询价", "本地视频缓存和项目 OSS 使用需分开核实；存储费用未算入 API 单价", "src/lib/server/aliyun-avatar-video.ts; services/oss.py", ""),
    ("egress", "流量", "视频下载/CDN/直播出站流量", "GB", "=B30", None, "待询价", "云端推流流量与云存储下载、CDN 回源按实际账单计", "src/lib/server/aliyun-avatar-video.ts; SRS", ""),
    ("license", "授权", "灵眸模板/形象/音色授权", "项", "=B31", None, "待合同价", "若已购且授权可复用则填 0；不要默认云形象免费", "src/lib/server/aliyun-avatar-video.ts", LINGMOU_GUIDE),
]

ASSUMPTIONS = [
    (5, "项目预算汇率", 7.25, "CNY/USD；仅预算换算，按付款账单复核"),
    (6, "每张图片输入 token", 2000, "含参考图，实际用量可能更高"),
    (7, "每张图片文本/思考输出 token", 1000, "示例值；按 Google usageMetadata 核对"),
    (8, "每张 1K 图片输出 token", 1120, "参考项目记录的典型 1K 用量"),
    (9, "Gemini 输入 USD/百万 token", 0.50, "参考项目成本表；当日官方页在本环境不可访问"),
    (10, "Gemini 文本/思考输出 USD/百万 token", 3.00, "参考项目成本表；当日官方页在本环境不可访问"),
    (11, "Gemini 图片输出 USD/百万 token", 60.00, "参考项目成本表；当日官方页在本环境不可访问"),
    (12, "HappyHorse 输入平均秒/条", 8, "请用成功任务 usage.duration 校准"),
    (13, "HappyHorse 输出平均秒/条", 8, "两者之和计费，非仅输出秒"),
    (14, "HappyHorse 成功条数/月", 50, "示例；与 Seedance 条数不重叠"),
    (15, "Seedance 成功条数/月", 20, "后备或切换生成，未并入 HappyHorse"),
    (16, "Seedance 输出平均秒/条", 30, "接口自动时长 -1；项目计价器缺少实际时长时暂按 30 秒"),
    (17, "灵眸成片条数/月", 100, "灵眸另行询价，不含已购形象授权"),
    (18, "灵眸平均成片秒/条", 60, "成片分钟只作预算量纲，合同计费方式待确认"),
    (19, "Gemini 图片张数/月", 200, "生成成功的 1K 图"),
    (20, "Qwen VC 合成字数/月", 100000, "云端语音合成，非灵眸内部已含 TTS"),
    (21, "Qwen VD 声音设计次数/月", 20, "设计音色的文字输入"),
    (22, "Qwen VD 平均字数/次", 200, "按真实账单字符数更新"),
    (23, "默认 GPT 调用次数/月", 5000, "脚本扩写/问答；没有代理单价所以不进入小计"),
    (24, "GPT 输入 token/次", 2000, "系统提示与图片输入可能增加 token"),
    (25, "GPT 输出 token/次", 500, "包含实际账单中的可计费输出"),
    (26, "Azure TTS 字数/月", 0, "默认 0，启用 Azure 时填写，勿与 Qwen 同批重复计数"),
    (27, "MuseTalk GPU 节点数/月", 1, "自有 GPU 时也有设备/电力成本；本表不虚构云租价"),
    (28, "业务主机数/月", 1, "已有服务器可填 0，新增云机按实报价"),
    (29, "对象存储 GB 月", 100, "按实际保存的视频时长与留存期调整"),
    (30, "公网出站 GB 月", 500, "CDN/直播平台按具体线路结算"),
    (31, "形象模板授权数", 1, "已付费且无需新购时填 0"),
    (32, "Fish Audio 查询次数/月", 0, "仅展示预设试听；无云合成基础用量"),
    (33, "Azure Neural USD/百万字符", 15.0, "Azure 官方定价页参考值；地区、免费额和合同折扣另核实"),
    (35, "Seedance 项目估价 CNY/百万 token", 69.91, "项目自估值，非 Seedance 官方合同价"),
    (36, "Seedance 输出宽度", 1280, "项目估算分辨率 720p"),
    (37, "Seedance 输出高度", 720, "项目估算分辨率 720p"),
    (38, "Seedance 输出 fps", 24, "按实际生成视频核对"),
]


def title(ws, name, note, columns):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.cell(1, 1, name)
    ws.cell(1, 1).font = Font(name="Microsoft YaHei", color=WHITE, bold=True, size=16)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=DARK)
    ws.row_dimensions[1].height = 39
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    ws.cell(2, 1, note)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 43


def header(ws, labels, row=4):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row, col, label)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(name="Microsoft YaHei", color=WHITE, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 34


def line(ws, row, values, unpriced=False):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        cell.font = Font(name="Microsoft YaHei", size=10, color=DARK)
        cell.fill = PatternFill("solid", fgColor=AMBER if unpriced else (LIGHT if row % 2 == 0 else WHITE))
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 52


def widths(ws, values):
    for col, width in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def build():
    wb = Workbook()
    overview = wb.active
    overview.title = "月费总览"
    detail = wb.create_sheet("云API及配套明细")
    inputs = wb.create_sheet("用量假设")
    evidence = wb.create_sheet("计价依据与边界")

    title(inputs, "每月调用量及计价输入", "B 列为可修改示例值；未取得合同单价的服务保持空白，不能把已知小计视为全量费用。", 3)
    header(inputs, ["假设项目", "值", "说明"])
    for row, label, value, note in ASSUMPTIONS:
        line(inputs, row, (label, value, note))
        inputs.cell(row, 2).number_format = '#,##0.00'
    widths(inputs, [43, 22, 80])
    inputs.freeze_panes = "B5"

    title(detail, "云端 API + 运行配套 | 月度报价清单", "按项目真实调用路径列项；示例金额是用量预算，不是账单。黄色行无可核验的价格，留空等待合同/云账单。", 11)
    header(detail, ["序号", "类别", "服务 / 模型", "计费单位", "每月用量", "单价 CNY", "月度小计 CNY", "价格状态", "计价口径及风险", "项目代码/成本表", "供应商资料"])
    for idx, (key, group, model, unit, qty, price, status, note, source, url) in enumerate(LINES, 1):
        row = 4 + idx
        quantity = f"='用量假设'!{qty[1:]}" if qty.startswith("=B") and qty[2:].isdigit() else qty.replace("B", "'用量假设'!B")
        unit_price = price.replace("B", "'用量假设'!B") if isinstance(price, str) else price
        line(detail, row, (idx, group, model, unit, quantity, unit_price,
                           f'=IF(OR(E{row}="",F{row}=""),"",E{row}*F{row})',
                           status, note, source, url), unpriced=price is None)
        for col, fmt in ((5, '#,##0.00'), (6, PRICE), (7, MONEY)):
            detail.cell(row, col).number_format = fmt
        if url:
            detail.cell(row, 11).hyperlink = url
            detail.cell(row, 11).style = "Hyperlink"
    last = 4 + len(LINES)
    total = last + 1
    line(detail, total, ("", "", "已知/暂估项目小计（不是完整月费）", "", "", "", f"=SUM(G5:G{last})", "未计价项另算"))
    detail.cell(total, 7).number_format = MONEY
    widths(detail, [8, 13, 56, 22, 17, 20, 20, 18, 82, 85, 83])
    detail.auto_filter.ref = f"A4:K{last}"
    detail.freeze_panes = "E5"

    title(overview, "AvatarLive | 现用云 API 月度费用预算", f"{DAY.isoformat()}  |  继续使用现有云 API + 本地 MuseTalk 实时渲染；无 H100 购置，仍需运行 GPU/业务服务。", 5)
    widths(overview, [30, 22, 26, 35, 115])
    header(overview, ["预算项目", "金额/数量", "单位", "是否完整", "说明"])
    line(overview, 5, ("已知/暂估费用小计", f"='云API及配套明细'!G{total}", "元/月", "否", "仅合计有单价项；含 Seedance 项目估值和 Gemini 按样本 token 的估值。"))
    overview.cell(5, 2).number_format = MONEY
    line(overview, 6, ("有用量但无价格", f'=COUNTIFS(\'云API及配套明细\'!E5:E{last},">0",\'云API及配套明细\'!F5:F{last},"")', "项", "必须另询", "灵眸、默认 GPT 代理、MuseTalk GPU、服务器/存储/流量/授权等未计价，不计作免费。"), True)
    baseline = 50 * (8 + 8) * 0.9 + 20 * 30 * (69.91 * 1280 * 720 * 24 / 1024 / 1_000_000)
    baseline += 200 * ((2000 * .5 + 1000 * 3 + 1120 * 60) / 1_000_000 * 7.25) + 100000 / 10000 * .8 + 20 * 200 / 10000 * .8
    line(overview, 8, ("静态预览（示例输入）", round(baseline, 2), "元/月", "否", "不随 Excel 输入变动；修改后以上方公式为准；这不是总月费。"))
    overview.cell(8, 2).number_format = MONEY
    notes = [
        (10, "视频链路", "灵眸负责透明数字人口播成片；HappyHorse 为动态底片主模型，Seedance 2.5 为独立的后备/切换任务；不能把同一任务重复计数。"),
        (11, "实时直播", "MuseTalk 在本地 GPU 渲染嘴型，云 API 并不能消除 GPU 需求；若改云 GPU，费用需据机型和时长单独询价。"),
        (12, "语音与文本", "Qwen VC 和 Azure 是不同路径；Azure 默认用量为零。GPT 是默认 llm-gpt 经内部 LiteLLM 转发，代理价无公开账单佐证。"),
        (13, "报价缺口", "灵眸合同单价、LLM 代理输入/输出单价、模板授权、GPU/业务主机及 OSS/流量等无验证价格；必须补齐后才有完整月费。"),
        (14, "计价差异", "Qwen3-TTS-VC 官网页展示 0.8 元/万字符，参考项目内部计算仍为 1.0；使用前请以账户账单核对并修订。"),
        (15, "价格属性", "已知项只反映公开价或项目内部估值；人民币预算汇率 7.25，未扣免费额、折扣，未含税、失败重试和运维人力。"),
    ]
    for row, label, value in notes:
        overview.cell(row, 1, label)
        overview.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT)
        overview.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        overview.cell(row, 2, value).alignment = Alignment(wrap_text=True, vertical="center")
        overview.row_dimensions[row].height = 43

    title(evidence, "价格依据、可选模型与待询事项", "公开网页/项目内部成本表只用于立项估算；账户地域、合同、模型变更及账单为最终口径。", 4)
    header(evidence, ["项目", "依据类型", "核算结论", "来源链接/项目文件"])
    evidence_rows = [
        ("灵眸数字人", "需供应商报价", "项目调用 CreateBroadcastVideoFromTemplate，1080p/30fps；未取得公开可核对的模板/口播价格。", LINGMOU_GUIDE),
        ("HappyHorse 1.0 Edit", "阿里云公开原价", "720P 0.90 元/秒，按输入+输出秒计；成功任务以 usage.duration 为准。", ALIYUN_PRICE),
        ("Seedance 2.5", "项目内部折算值", "69.91 元/百万 token；720p/24fps 约 1.510056 元/输出秒，不是官方合同报价。", "seo_video_generate/services/cost/seedance_cost.py"),
        ("Gemini 3.1 Flash Image", "项目内部计价器", "输入 $0.50/M、文本输出 $3/M、图片输出 $60/M；参考页本环境连接超时，按 usage 核实。", GOOGLE_PRICE),
        ("Qwen3-TTS VC/VD", "阿里云公开页", "VC 与 VD 页显示 0.80 元/万字符；VC 项目旧计算器仍为 1.00，可能导致统计差异。", ALIYUN_PRICE),
        ("Azure Neural TTS", "Azure 定价页", "页面 Neural/Neural HD Flash 显示 $15/百万字符，按 7.25 预算换算；账户地域/折扣需核实。", AZURE_PRICE),
        ("默认 GPT 5.4", "内部代理待询", "项目默认 azure-gpt-5.4 via LiteLLM；OpenAI 官方文档本环境 403，且代理结算不等于公开单价。", "AvatarLive/apps/api/app/services/llm/config.py"),
        ("可选 LLM", "不与默认叠加", "deepseek-v4-pro、gemini-3-pro-preview、doubao-seed-2-0-pro-260215 可切换；需对应代理价及用量。", "AvatarLive/apps/api/app/services/llm/config.py"),
        ("本地运行依赖", "另行预算", "MuseTalk/OpenVoice 仍是本地推理；云 GPU、CPU、OSS/CDN、形象授权等尚未取得单价。", "AvatarLive/apps/musetalk/server.py"),
    ]
    for row, item in enumerate(evidence_rows, 5):
        line(evidence, row, item)
        if item[3].startswith("https://"):
            evidence.cell(row, 4).hyperlink = item[3]
            evidence.cell(row, 4).style = "Hyperlink"
    widths(evidence, [32, 25, 115, 100])
    evidence.freeze_panes = "B5"
    for ws in wb:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"{OUTPUT}\nline_items={len(LINES)} example_known_subtotal_cny={baseline:.2f}")


if __name__ == "__main__":
    build()
